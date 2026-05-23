# -*- coding: utf-8 -*-
"""latest_template_builder.py

把【出货资料_YYYYMMDD.xlsx】里的数据，按你的映射规则生成【最新模板_YYYYMMDD.xlsx】。

- 仅用于【监控推送 + 闲鱼】订单（由 app.py 决定是否调用）
- 不依赖任何“底稿/参考文件”，直接用 openpyxl 从零创建 Excel

输出模板结构（与您提供的“最新模板.xlsx”一致）：
- sheet: 线上贴单资料
  表头：订单编号, 店铺名, 收件人, 配送方式, 货物标题, 货物品名, 货物规格, 快递单号, 报单备注, 货单备注
- sheet: 宅配打包资料
  表头：订单编号, 店铺名, 收件人, 收件人电话, 收件人地址, 配送方式, 货物标题, 货物品名, 货物规格, 价格, 快递单号, 报单备注, 货单备注
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List

import re

import openpyxl


def _norm(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        try:
            if float(v).is_integer():
                return str(int(v))
        except Exception:
            pass
        return str(v)
    return str(v).strip()


def _merge_name_spec(items: List[tuple]) -> tuple:
    """合并多行的(名称, 规格)。

    规则：按(名称, 单位)分组，同组数量累加，不同组用+拼接。
    例：
      [(擺件,5個),(擺件,5個),(擺件,5個)] → (擺件, 15個)
      [(擺件,5個),(項鏈,1條)]            → (擺件+項鏈, 5個+1條)
      [(擺件,1個),(擺件,1雙)]            → (擺件+擺件, 1個+1雙)
    """
    if not items:
        return "", ""

    from collections import OrderedDict
    groups: OrderedDict = OrderedDict()  # key=(name, unit) → total_qty
    for name, spec in items:
        name = (name or "").strip()
        spec = (spec or "").strip()
        m = re.match(r"^(\d+)\s*(.+)$", spec)
        if m:
            qty = int(m.group(1))
            unit = m.group(2).strip()
        else:
            qty = 0
            unit = spec
        key = (name, unit)
        if key in groups:
            groups[key] += qty
        else:
            groups[key] = qty

    names = []
    specs = []
    for (name, unit), total_qty in groups.items():
        names.append(name)
        specs.append(f"{total_qty}{unit}" if total_qty > 0 else unit)

    return "+".join(names), "+".join(specs)


def _sheet_rows_as_dicts(wb: openpyxl.Workbook, sheet_name: str) -> List[Dict[str, Any]]:
    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    try:
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    except Exception:
        return []
    headers = [_norm(x) for x in header_row]
    # 去掉尾部空表头
    while headers and (not headers[-1]):
        headers.pop()
    rows: List[Dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        d: Dict[str, Any] = {}
        empty = True
        for i, h in enumerate(headers):
            if not h:
                continue
            v = row[i] if i < len(row) else None
            if v not in (None, ""):
                empty = False
            d[h] = v
        if empty:
            continue
        rows.append(d)
    return rows


def _query_d1_voided_pks(log_fn=None, *, _raise: bool = True) -> set:
    """v6.0.68 加 / v6.0.80 改:從 D1 撈過去 30 天作廢的業績 PK(perf_code|order_code)。

    回 set,失敗時:
      - _raise=True (默認,SYB 場景用):直接拋異常 → SYB 看到後設 _voided_pks=None (保守不 +N)
      - _raise=False (builder 場景用):靜默回空 set,builder 退回不過濾的行為

    用於 SYB 上傳判斷作廢、模板 builder 過濾作廢 row。

    v6.1.45:帶 owner=_get_my_owner() — D1 worker 對員工強制要求 owner,
    沒帶會 reject「员工查询必须指定 owner」,導致 SYB 撞「已存在」時誤跳過。
    修「自動出貨第二次上傳被 SYB 誤判跳過」bug
    """
    from datetime import datetime, timedelta
    from .performance_feature import query_records, is_voided_note, _get_my_owner
    today = datetime.now().strftime("%Y-%m-%d")
    from_d = (datetime.now() - timedelta(days=30)).strftime("%Y-%m-%d")
    # v6.1.45:帶 owner 避開 worker「员工必須指定 owner」權限檢查
    _owner = _get_my_owner() or ""
    try:
        # v6.0.80:raise_on_fail=True,失敗時拋異常給上層,不靜默回空 list
        rows = query_records(owner=_owner, from_date=from_d, to_date=today, limit=5000, raise_on_fail=True)
    except Exception as e:
        if log_fn:
            log_fn(f"[builder] ⚠️ D1 query 作廢清單失敗: {e}")
        if _raise:
            raise  # SYB 場景:讓上層接到後 _voided_pks=None 保守處理
        return set()  # builder 場景:靜默 fallback

    voided = set()
    for r in rows:
        if is_voided_note(r.get("note", "")):
            pk = r.get("pk", "")
            if pk:
                voided.add(pk)
    if log_fn:
        log_fn(f"[builder] D1 voided pks (近 30 天): {len(voided)} 筆 / 總共拉到 {len(rows)} 筆業績")
    return voided


def build_latest_template_from_ship_excel(
    ship_excel_path: Path,
    output_path: Path,
    order_nos: List[str],
    fixed_note: str = "鐘鐘件鐘鐘件",
    log_fn=None,
) -> Path:
    """从【出货资料】生成【最新模板】。

    Args:
        ship_excel_path: 出货资料_YYYYMMDD.xlsx
        output_path:     最新模板_YYYYMMDD.xlsx
        order_nos:       只生成这些“主订单号”的行
        fixed_note:      报单备注固定值（默认：鐘鐘件鐘鐘件）
        log_fn:          日誌函數(可選),用於記錄 D1 過濾、dedup 動作

    Returns:
        output_path
    """
    ship_excel_path = Path(ship_excel_path)
    output_path = Path(output_path)

    if not ship_excel_path.exists():
        raise FileNotFoundError(f"ship excel not found: {ship_excel_path}")

    need = {str(x).strip() for x in (order_nos or []) if str(x).strip()}
    if not need:
        # 没有需要生成的订单，直接不产出
        raise ValueError("order_nos is empty")

    # v6.0.68 ★:撈 D1 作廢清單,後面用來過濾「重綁前舊 perf_code」的 row
    # v6.0.80:builder 場景失敗時靜默 fallback(不過濾),不影響模板生成
    _voided_pks = _query_d1_voided_pks(log_fn=log_fn, _raise=False)

    wb = openpyxl.load_workbook(ship_excel_path, data_only=True)

    online_rows = _sheet_rows_as_dicts(wb, "线上贴单资料")
    home_rows = _sheet_rows_as_dicts(wb, "宅配打包资料")

    # 过滤只保留需要的主订单
    def _get_main_code(s: str) -> str:
        s = (s or "").strip()
        if not s:
            return ""
        # 兼容 123+234 / 123＋234
        for sep in ("+", "＋"):
            if sep in s:
                s = s.split(sep, 1)[0].strip()
        return s

    # v6.0.68 ★:過濾條件 = 訂單號在 need 內 AND (perf_code|order_code) 不在作廢清單
    # 續行(編碼空)隨主行一起被決定保留與否
    def _is_voided(perf_code: str, order_code: str) -> bool:
        if not _voided_pks or not perf_code or not order_code:
            return False
        return f"{perf_code}|{order_code}" in _voided_pks

    online_pick: List[Dict[str, Any]] = []
    _last_online_main = ""
    _last_online_perf = ""
    _last_was_voided = False  # 主行作廢 → 續行也跳過
    _voided_count = 0
    for d in online_rows:
        main = _get_main_code(_norm(d.get("系统編碼") or d.get("系统编码") or ""))
        if main:
            _last_online_main = main
            cur_perf = _norm(d.get("編碼") or d.get("编码") or "")
            if cur_perf:
                _last_online_perf = cur_perf
            _last_was_voided = _is_voided(_last_online_perf, _last_online_main)
            if _last_was_voided:
                _voided_count += 1
                if log_fn:
                    log_fn(f"[builder] 跳過作廢的 row: perf={_last_online_perf} order={_last_online_main}")
        else:
            # 续行（同一订单的第2、3...个快递）：继承上一行的主订单号 + 主行作廢狀態
            main = _last_online_main
        if not main or main not in need:
            continue
        if _last_was_voided:
            continue  # 作廢的主行 + 它的續行都跳過
        online_pick.append(d)

    if log_fn and _voided_count:
        log_fn(f"[builder] 線上貼單:過濾掉 {_voided_count} 個作廢主行(連同續行)")

    home_pick: List[Dict[str, Any]] = []
    _last_home_main = ""
    _last_home_perf = ""
    _last_was_voided_h = False
    _voided_count_h = 0
    for d in home_rows:
        main = _get_main_code(_norm(d.get("訂單編碼") or d.get("订单编码") or ""))
        if main:
            _last_home_main = main
            cur_perf = _norm(d.get("編碼") or d.get("编码") or "")
            if cur_perf:
                _last_home_perf = cur_perf
            _last_was_voided_h = _is_voided(_last_home_perf, _last_home_main)
            if _last_was_voided_h:
                _voided_count_h += 1
                if log_fn:
                    log_fn(f"[builder] 跳過作廢的 row: perf={_last_home_perf} order={_last_home_main}")
        else:
            main = _last_home_main
        if not main or main not in need:
            continue
        if _last_was_voided_h:
            continue
        home_pick.append(d)

    if log_fn and _voided_count_h:
        log_fn(f"[builder] 宅配:過濾掉 {_voided_count_h} 個作廢主行(連同續行)")

    # 如果模板已存在,加载已有数据(追加模式,避免覆盖之前的记录)
    # v6.0.68 ★:載入時記下 (order, tracking, shop) gkey 用於後面去重
    # 預期欄位順序(最新模板格式):
    #   线上:  订单编号(0), 店铺名(1), 收件人(2), 配送方式(3), 货物标题(4),
    #         货物品名(5), 货物规格(6), 快递单号(7), 报单备注(8), 货单备注(9)
    #   宅配:  订单编号(0), 店铺名(1), 收件人(2), 收件人电话(3), 收件人地址(4),
    #         配送方式(5), 货物标题(6), 货物品名(7), 货物规格(8), 价格(9),
    #         快递单号(10), 报单备注(11), 货单备注(12)
    existing_online_with_keys: List[tuple] = []  # [(gkey, row_list), ...]
    existing_home_with_keys: List[tuple] = []
    if output_path.exists():
        try:
            _ewb = openpyxl.load_workbook(output_path, data_only=True)
            if "线上贴单资料" in _ewb.sheetnames:
                _ws = _ewb["线上贴单资料"]
                for _row in _ws.iter_rows(min_row=2, values_only=True):
                    if _row and any(c not in (None, "") for c in _row):
                        rl = list(_row)
                        gkey = (
                            _get_main_code(_norm(rl[0] if len(rl) > 0 else "")),  # order_no
                            _norm(rl[7] if len(rl) > 7 else ""),                  # tracking
                            _norm(rl[1] if len(rl) > 1 else ""),                  # shop
                        )
                        existing_online_with_keys.append((gkey, rl))
            if "宅配打包资料" in _ewb.sheetnames:
                _ws = _ewb["宅配打包资料"]
                for _row in _ws.iter_rows(min_row=2, values_only=True):
                    if _row and any(c not in (None, "") for c in _row):
                        rl = list(_row)
                        gkey = (
                            _get_main_code(_norm(rl[0] if len(rl) > 0 else "")),  # order_no
                            _norm(rl[10] if len(rl) > 10 else ""),                # tracking
                            _norm(rl[1] if len(rl) > 1 else ""),                  # shop
                        )
                        existing_home_with_keys.append((gkey, rl))
            _ewb.close()
        except Exception as e:
            if log_fn:
                log_fn(f"[builder] ⚠️ 讀取既有模板失敗: {e}")

    # 创建输出 workbook
    out_wb = openpyxl.Workbook()
    # 删除默认 sheet
    try:
        default = out_wb.active
        out_wb.remove(default)
    except Exception:
        pass

    # 线上贴单资料
    headers_online = [
        "订单编号",
        "店铺名",
        "收件人",
        "配送方式",
        "货物标题",
        "货物品名",
        "货物规格",
        "快递单号",
        "报单备注",
        "货单备注",
    ]
    ws_o = out_wb.create_sheet("线上贴单资料")
    ws_o.append(headers_online)

    # 注:existing rows 的寫入挪到 new groups 算完 之後,因為要先知道 new groups 的 gkey 才能去重
    # (見 v6.0.68 修補:existing 跟 new groups 撞 gkey 時新算的贏)

    # 记住主行字段,续行(系统編碼为空)继承主行的订单信息
    # 先按(主订单号+快递单号+编码)收集所有行,再合并名称/规格
    # ⚠ 加 shop(perf_code) 进 group key:同订单同 tracking 但不同 perf_code 必须分组
    #   场景:出货失败重做 → row2 perf=白050401, row3 perf=白050402,tracking 相同
    #         不加 shop 时第二条 perf_code 会被丢失,模板上两行都变白050401
    _prev_online: Dict[str, str] = {}
    _online_groups: Dict[tuple, Dict] = {}  # key=(order_no, tracking, shop) → {字段 + items:[(name,spec)]}
    for d in online_pick:
        order_no = _norm(d.get("系统編碼") or d.get("系统编码") or "")
        shop = _norm(d.get("編碼") or d.get("编码") or "")
        receiver = _norm(d.get("收件人") or "")
        channel = _norm(d.get("渠道") or "")
        name = _norm(d.get("商品名稱") or d.get("商品名称") or "")
        spec = _norm(d.get("规格") or d.get("規格") or "")
        tracking = _norm(d.get("国内快递单号") or d.get("国内快递單號") or "")
        remark = _norm(d.get("备注") or d.get("備注") or "")

        if order_no:
            _prev_online = {"order_no": order_no, "shop": shop, "receiver": receiver,
                            "channel": channel, "name": name, "spec": spec, "remark": remark}
        else:
            order_no = _prev_online.get("order_no", "")
            shop = shop or _prev_online.get("shop", "")
            receiver = receiver or _prev_online.get("receiver", "")
            channel = channel or _prev_online.get("channel", "")
            if not name:
                name = _prev_online.get("name", "")
            if not spec:
                spec = _prev_online.get("spec", "")
            if not remark:
                remark = _prev_online.get("remark", "")

        order_no = _get_main_code(order_no)
        gkey = (order_no, tracking, shop)
        if gkey not in _online_groups:
            _online_groups[gkey] = {"shop": shop, "receiver": receiver,
                                    "channel": channel, "remark": remark, "items": []}
        _online_groups[gkey]["items"].append((name, spec))

    # v6.0.68 ★:先寫 existing(過濾掉:1)跟 new groups 撞 gkey 的 2)existing 自己內部 dup 3)作廢的 perf_code)
    _seen_existing_gkeys = set()
    _kept_existing = 0
    _dropped_collide = 0
    _dropped_dup = 0
    _dropped_voided = 0
    for gkey, rl in existing_online_with_keys:
        # 跟 new groups 撞 → 新算的贏,丟 existing
        if gkey in _online_groups:
            _dropped_collide += 1
            continue
        # existing 自己內部 dup
        if gkey in _seen_existing_gkeys:
            _dropped_dup += 1
            continue
        # existing 是已作廢的 perf_code(舊 bug 殘留或跨日作廢) → 丟
        if _voided_pks and gkey[0] and gkey[2]:
            if f"{gkey[2]}|{gkey[0]}" in _voided_pks:
                _dropped_voided += 1
                continue
        _seen_existing_gkeys.add(gkey)
        ws_o.append(rl)
        _kept_existing += 1
    if log_fn and (existing_online_with_keys):
        log_fn(f"[builder] 線上 existing 處理:保留 {_kept_existing}, "
               f"撞新算丟 {_dropped_collide}, 內部 dup 丟 {_dropped_dup}, 作廢丟 {_dropped_voided}")

    for (order_no, tracking, _shop), g in _online_groups.items():
        merged_name, merged_spec = _merge_name_spec(g["items"])
        ws_o.append([
            order_no,
            g["shop"],
            g["receiver"],
            g["channel"],
            merged_name,
            merged_name,
            merged_spec,
            tracking,
            fixed_note,
            g["remark"],
        ])

    # 宅配打包资料
    headers_home = [
        "订单编号",
        "店铺名",
        "收件人",
        "收件人电话",
        "收件人地址",
        "配送方式",
        "货物标题",
        "货物品名",
        "货物规格",
        "价格",
        "快递单号",
        "报单备注",
        "货单备注",
    ]
    ws_h = out_wb.create_sheet("宅配打包资料")
    ws_h.append(headers_home)

    # 注:existing rows 寫入挪到 new groups 算完之後(同上)

    # 记住主行字段,续行继承主行的订单信息
    # group key 同上(线上贴单),加 shop 区分「重做」场景的不同 perf_code
    _prev_home: Dict[str, str] = {}
    _home_groups: Dict[tuple, Dict] = {}
    for d in home_pick:
        order_no = _norm(d.get("訂單編碼") or d.get("订单编码") or "")
        shop = _norm(d.get("編碼") or d.get("编码") or "")
        receiver = _norm(d.get("收件人") or "")
        phone = _norm(d.get("联系电话") or d.get("聯係電話") or d.get("联係电话") or d.get("联繫电话") or "")
        addr = _norm(d.get("收件地址") or "")
        name = _norm(d.get("商品名稱") or d.get("商品名称") or "")
        spec = _norm(d.get("規格") or d.get("规格") or "")
        tracking = _norm(d.get("国内快递单号") or d.get("国内快递單號") or "")
        remark = _norm(d.get("備注") or d.get("备注") or "")

        if order_no:
            _prev_home = {"order_no": order_no, "shop": shop, "receiver": receiver,
                          "phone": phone, "addr": addr, "name": name, "spec": spec, "remark": remark}
        else:
            order_no = _prev_home.get("order_no", "")
            shop = shop or _prev_home.get("shop", "")
            receiver = receiver or _prev_home.get("receiver", "")
            phone = phone or _prev_home.get("phone", "")
            addr = addr or _prev_home.get("addr", "")
            if not name:
                name = _prev_home.get("name", "")
            if not spec:
                spec = _prev_home.get("spec", "")
            if not remark:
                remark = _prev_home.get("remark", "")

        order_no = _get_main_code(order_no)
        gkey = (order_no, tracking, shop)
        if gkey not in _home_groups:
            _home_groups[gkey] = {"shop": shop, "receiver": receiver, "phone": phone,
                                  "addr": addr, "remark": remark, "items": []}
        _home_groups[gkey]["items"].append((name, spec))

    # v6.0.68 ★ 同線上邏輯:先寫 existing(過濾撞 + dup + 作廢),再寫 new groups
    _seen_existing_gkeys_h = set()
    _kept_existing_h = 0
    _dropped_collide_h = 0
    _dropped_dup_h = 0
    _dropped_voided_h = 0
    for gkey, rl in existing_home_with_keys:
        if gkey in _home_groups:
            _dropped_collide_h += 1
            continue
        if gkey in _seen_existing_gkeys_h:
            _dropped_dup_h += 1
            continue
        if _voided_pks and gkey[0] and gkey[2]:
            if f"{gkey[2]}|{gkey[0]}" in _voided_pks:
                _dropped_voided_h += 1
                continue
        _seen_existing_gkeys_h.add(gkey)
        ws_h.append(rl)
        _kept_existing_h += 1
    if log_fn and existing_home_with_keys:
        log_fn(f"[builder] 宅配 existing 處理:保留 {_kept_existing_h}, "
               f"撞新算丟 {_dropped_collide_h}, 內部 dup 丟 {_dropped_dup_h}, 作廢丟 {_dropped_voided_h}")

    for (order_no, tracking, _shop), g in _home_groups.items():
        merged_name, merged_spec = _merge_name_spec(g["items"])
        ws_h.append([
            order_no,
            g["shop"],
            g["receiver"],
            g["phone"],
            g["addr"],
            "黑貓",  # 固定
            merged_name,
            merged_name,
            merged_spec,
            0,  # 代收货款/价格：固定 0
            tracking,
            fixed_note,
            g["remark"],
        ])

    # 保存
    output_path.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(output_path)
    return output_path
